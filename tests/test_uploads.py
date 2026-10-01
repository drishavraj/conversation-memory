from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from sqlalchemy import select
from test_api import client, headers
from app.database import session_factory
from app.models import Asset, Entry, Job
from app.worker import process_one
from app.documents import parse_document, DocumentError
from test_processing import Fake, RESULT, SOURCE
import pytest


def upload(client,headers,name,content,**metadata):
    return client.post('/api/entries/upload',headers=headers,files={'file':(name,content)},data=metadata)

def test_txt_upload_process_and_download(client,headers):
    result=upload(client,headers,'call.txt',SOURCE.encode(),theme_id='office',event_at='2026-10-01T12:00:00+05:30')
    assert result.status_code==201
    id=result.json()['id']
    assert client.get('/api/entries/'+id+'/file',headers=headers).content==SOURCE.encode()
    assert client.get('/api/entries/'+id+'/file').status_code==401
    process_one(session_factory(client.app.state.engine),Fake())
    detail=client.get('/api/entries/'+id,headers=headers).json()
    assert detail['status']=='ready' and detail['original_text']==SOURCE
    assert upload(client,headers,'renamed.txt',SOURCE.encode(),theme_id='office',event_at='2026-10-01T06:30:00Z').json()['id']==id

def test_upload_validation(client,headers):
    assert upload(client,{},'call.txt',b'hello').status_code==401
    assert upload(client,headers,'call.exe',b'hello').status_code==415
    assert upload(client,headers,'call.pdf',b'not PDF').status_code==422
    assert upload(client,headers,'empty.txt',b'').status_code==422
    assert upload(client,headers,'call.txt',b'hello',theme_id='bad').status_code==422
    assert upload(client,headers,'call.txt',b'hello',event_at='2026-10-01T12:00:00').status_code==422
    assert upload(client,headers,'big.txt',b'x'*(10*1024*1024+1)).status_code==413

def test_docx_and_pdf_parsing():
    doc=Document();doc.add_paragraph('Synthetic discussion');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Task';table.cell(0,1).text='Send document'
    out=BytesIO();doc.save(out)
    text=parse_document(out.getvalue(),'.docx')
    assert 'Synthetic discussion' in text and 'Send document' in text
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);out=BytesIO();writer.write(out)
    with pytest.raises(DocumentError,match='ocr_not_supported'):parse_document(out.getvalue(),'.pdf')
    with pytest.raises(DocumentError):parse_document(b'bad','.docx')
    with pytest.raises(DocumentError):parse_document(b'\xff','.txt')

def test_audio_transcription_checkpoint_survives_retry(client,headers):
    wav=b'RIFF'+b'\0'*4+b'WAVE'+b'data'+b'\0'*8
    id=upload(client,headers,'call.wav',wav,theme_id='office').json()['id']
    class Audio(Fake):
        transcriptions=0
        fail=True
        def transcribe(self,*args):self.transcriptions+=1;return SOURCE
        def extract(self,*args):
            if self.fail:raise RuntimeError('failed')
            return super().extract(*args)
    provider=Audio();sessions=session_factory(client.app.state.engine)
    process_one(sessions,provider)
    detail=client.get('/api/entries/'+id,headers=headers).json()
    assert detail['status']=='failed' and detail['original_text']==SOURCE
    assert client.post('/api/entries/'+id+'/retry',headers=headers).status_code==200
    provider.fail=False;process_one(sessions,provider)
    assert provider.transcriptions==1
    assert client.get('/api/entries/'+id,headers=headers).json()['status']=='ready'

def test_text_based_pdf_extracts_content():
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 250 Td (Credentials are pending.) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    out=BytesIO();writer.write(out)
    assert 'Credentials are pending.' in parse_document(out.getvalue(),'.pdf')
