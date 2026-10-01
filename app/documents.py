from io import BytesIO
from zipfile import ZipFile, BadZipFile
from pypdf import PdfReader
from docx import Document

class DocumentError(ValueError):
    pass

MAX_TEXT=200000

def parse_document(content, suffix):
    try:
        if suffix==".txt":
            text=content.decode("utf-8-sig")
        elif suffix==".docx":
            with ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>30*1024*1024:
                    raise DocumentError("document_expanded_too_large")
            doc=Document(BytesIO(content))
            parts=[p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
            text="\n".join(parts)
        elif suffix==".pdf":
            reader=PdfReader(BytesIO(content))
            if reader.is_encrypted:raise DocumentError("encrypted_pdf_unsupported")
            if len(reader.pages)>100:raise DocumentError("pdf_too_many_pages")
            parts=[]
            size=0
            for page in reader.pages:
                part=page.extract_text() or ""
                size+=len(part)
                if size>MAX_TEXT:raise DocumentError("document_text_too_large")
                parts.append(part)
            text="\n\n".join(parts)
        else:raise DocumentError("unsupported_document")
    except DocumentError:raise
    except Exception:raise DocumentError("document_parse_failed") from None
    if not text.strip():raise DocumentError("no_text_found_ocr_not_supported")
    if len(text)>MAX_TEXT:raise DocumentError("document_text_too_large")
    if "\x00" in text:raise DocumentError("invalid_text")
    return text
